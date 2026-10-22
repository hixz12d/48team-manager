"""One-off import of the team purchase ledger from the Excel workbook (Pro管理_优化版.xlsx).

Two steps, because the workbook is on this machine and the database is in the VPS container:

  extract  (local, needs openpyxl; no app imports):
      py -3 scripts/import_purchase.py extract <xlsx> --out .tmp/purchase-import.json
  apply    (inside the container, from /app; dry run unless --commit):
      python /tmp/import_purchase.py apply /tmp/purchase-import.json [--commit] [--overwrite] [--assign 行号=团队ID]

The JSON holds 2FA secrets in plain text: keep it under .tmp/ (or /tmp in the container)
and delete it right after use. Output only ever says whether a 2FA exists.
Only the local ``workspaces`` table is written; official Team, Sub2API and codex-rs are untouched.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import re
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

SHEET = "账号"
HEADER_ROW = 7
IN_USE = "使用中"
# Column letters → 1-based index: B 账号, C 状态, E 到期日, H 支付渠道, I 费用, J 币种, L 2FA, P 备注
COLUMNS = {"account": 2, "status": 3, "expires_on": 5, "payment_card": 8, "cost": 9, "currency": 10, "totp_secret": 12, "note": 16}
EXPECTED_HEADERS = {"account": "账号", "status": "状态", "expires_on": "到期日", "payment_card": "支付渠道",
                    "cost": "费用", "currency": "币种", "totp_secret": "2FA", "note": "备注"}
EMAIL = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")
CURRENCY_ALIASES = {"USD": "USD", "US$": "USD", "$": "USD", "美元": "USD", "CNY": "CNY", "RMB": "CNY", "¥": "CNY", "￥": "CNY", "人民币": "CNY"}
CARD_BRANDS = ("kast", "mexc", "lava", "plasma")
GMAIL_SWAP = {"gmail.com": "googlemail.com", "googlemail.com": "gmail.com"}


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _cost(value) -> str | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return format(Decimal(str(value).strip()).normalize(), "f")
    except InvalidOperation:
        return str(value).strip()


def _date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def extract(xlsx: Path, out: Path) -> None:
    import openpyxl

    # data_only: 状态、币种等列是公式，取 Excel 缓存的计算结果。
    workbook = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    sheet = workbook[SHEET]
    header = next(sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
    for key, index in COLUMNS.items():
        if _text(header[index - 1]) != EXPECTED_HEADERS[key]:
            raise SystemExit(f"表头不对：第 {index} 列应为“{EXPECTED_HEADERS[key]}”，实际是“{header[index - 1]}”")
    items = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True), start=HEADER_ROW + 1):
        cell = lambda key: row[COLUMNS[key] - 1] if len(row) >= COLUMNS[key] else None
        account = _text(cell("account"))
        if not account or _text(cell("status")) != IN_USE:
            continue
        account = account.lower()
        currency_raw = _text(cell("currency"))
        items.append({
            "row": row_number,
            "account": account,
            "is_email": bool(EMAIL.fullmatch(account)),
            "payment_card": _text(cell("payment_card")),
            "cost": _cost(cell("cost")),
            "currency": CURRENCY_ALIASES.get((currency_raw or "").upper(), currency_raw),
            "expires_on": _date(cell("expires_on")),
            "totp_secret": _text(cell("totp_secret")),
            "note": _text(cell("note")),
            "workspace_id": None,
        })
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({"source": str(xlsx), "items": items}, ensure_ascii=False, indent=1), encoding="utf-8")
    not_email = [item["account"] for item in items if not item["is_email"]]
    print(f"已读取 {len(items)} 行“{IN_USE}”，写入 {out}（含 2FA 明文，用完请删除）")
    print(f"2FA 有 {sum(bool(item['totp_secret']) for item in items)} 行")
    print("不是邮箱的账号：" + ("、".join(not_email) if not_email else "无"))


def _card_core(text: str | None) -> str:
    # “Kast银色卡2”“Kast银色 10061”都归成“kast银色”；MECX 是团队名里的常见笔误。
    return re.sub(r"[\s\d卡]+", "", str(text or "")).lower().replace("mecx", "mexc")


def _card_mismatch(card: str | None, team_name: str | None) -> bool:
    """只在团队名里写了卡名（含 Kast / MEXC 等）且与 Excel 支付卡对不上时提示。"""
    name = _card_core(team_name)
    if not card or not any(brand in name for brand in CARD_BRANDS):
        return False
    return _card_core(card) != name


async def apply(path: Path, commit: bool, overwrite: bool, assign: dict[int, int] | None = None) -> None:
    sys.path[:0] = [str(Path.cwd()), str(Path(__file__).resolve().parents[1])]
    from sqlalchemy import func, select, text

    from app.application.workspace_expiry import update_workspace_expiry
    from app.application.workspace_purchase import (
        normalize_cost, normalize_currency, normalize_totp_secret, purchase_record, update_workspace_purchase,
    )
    from app.core.config import Settings
    from app.persistence.database import create_engine, create_session_factory
    from app.persistence.models.identity import Account, Workspace

    items = json.loads(path.read_text(encoding="utf-8"))["items"]
    for item in items:
        if item["row"] in (assign or {}):
            item["workspace_id"] = assign[item["row"]]
    engine = create_engine(Settings())
    try:
        async with engine.connect() as conn:
            columns = {row[1] for row in (await conn.execute(text("PRAGMA table_info(workspaces)"))).fetchall()}
        if "owner_totp_secret_encrypted" not in columns:
            raise SystemExit("数据库还没有购买信息列：先部署包含购买信息的新版本（启动时自动补列），再运行导入。")
        session_factory = create_session_factory(engine)
        async with session_factory() as db:
            plans = []
            for item in items:
                issues, blocking, swapped = [], False, False
                workspaces: list = []
                if item.get("workspace_id"):
                    workspace = await db.get(Workspace, int(item["workspace_id"]))
                    workspaces = [workspace] if workspace else []
                    if not workspaces:
                        issues.append(f"workspace_id {item['workspace_id']} 不存在")
                elif item["is_email"]:
                    candidates = [item["account"]]
                    local, _, domain = item["account"].rpartition("@")
                    if domain in GMAIL_SWAP:
                        candidates.append(f"{local}@{GMAIL_SWAP[domain]}")
                    for index, email in enumerate(candidates):
                        account = await db.scalar(select(Account).where(func.lower(Account.email) == email))
                        if account is None:
                            continue
                        workspaces = list(await db.scalars(select(Workspace).where(Workspace.owner_account_id == account.id).order_by(Workspace.id)))
                        if workspaces and index:
                            swapped = True
                            issues.append(f"按 {email} 匹配（gmail / googlemail 互换）")
                        if workspaces and account.local_purpose != "mother":
                            issues.append(f"该账号本地用途是 {account.local_purpose}，不是母号")
                        if workspaces:
                            break
                    if not workspaces:
                        issues.append("未匹配：邮箱找不到所属团队，请在 JSON 里填 workspace_id")
                else:
                    issues.append("未匹配：不是邮箱，请在 JSON 里填 workspace_id")
                if len(workspaces) > 1:
                    issues.append("一个邮箱对应多个团队：" + "、".join(f"#{w.id}" for w in workspaces) + "，请在 JSON 里填 workspace_id")
                    blocking = True
                workspace = workspaces[0] if len(workspaces) == 1 else None
                fields = {}
                for key in ("payment_card", "note"):
                    if item.get(key):
                        fields[key] = item[key]
                try:
                    if item.get("cost") is not None:
                        fields["cost"] = normalize_cost(item["cost"])
                    if item.get("currency"):
                        fields["currency"] = normalize_currency(item["currency"])
                except ValueError as exc:
                    issues.append(f"费用 / 币种无效：{exc}")
                    blocking = True
                if item.get("totp_secret"):
                    try:
                        fields["totp_secret"] = normalize_totp_secret(item["totp_secret"])
                    except ValueError:
                        issues.append("2FA 格式不对，这一项不导入")
                db_expiry = None
                if workspace is not None:
                    team_name = workspace.custom_name or workspace.name or workspace.official_name
                    if _card_mismatch(item.get("payment_card"), team_name):
                        issues.append("支付卡与团队名里的卡名不一致（以 Excel 为准写支付卡，不改团队名）")
                    db_expiry = workspace.manual_expires_on.isoformat() if workspace.manual_expires_on else None
                    if item.get("expires_on") and db_expiry and db_expiry != item["expires_on"]:
                        issues.append(f"到期日不同：库里 {db_expiry}，保留库里")
                    existing = purchase_record(workspace)
                    if any(existing[key] for key in ("payment_card", "cost", "note")) or existing["totp_state"] == "stored":
                        issues.append("团队已有购买信息" + ("，--overwrite 覆盖" if overwrite else "，跳过（加 --overwrite 才覆盖）"))
                        blocking = blocking or not overwrite
                plans.append({"item": item, "workspace": workspace, "fields": fields, "issues": issues,
                              "blocking": blocking, "db_expiry": db_expiry, "swapped": swapped})

            seen: dict[int, list] = {}
            for plan in plans:
                if plan["workspace"] is not None:
                    seen.setdefault(plan["workspace"].id, []).append(plan)
            for same in seen.values():
                # 邮箱完全一致的优先；只靠 gmail / googlemail 互换匹配上的让出这个团队，算未匹配。
                exact = [p for p in same if not p["swapped"]]
                if len(same) > 1 and len(exact) == 1:
                    for plan in same:
                        if plan["swapped"]:
                            plan["issues"].append(f"该团队已由第 {exact[0]['item']['row']} 行（邮箱完全一致）占用，本行算未匹配")
                            plan["workspace"], plan["blocking"] = None, True
                    continue
                if len(same) > 1:
                    rows = "、".join(str(p["item"]["row"]) for p in same)
                    for plan in same:
                        plan["issues"].append(f"Excel 第 {rows} 行对应同一个团队，都不写入")
                        plan["blocking"] = True

            print("行 | 账号 | 团队 | 支付卡 | 费用 | 币种 | 到期（Excel / 库里）| 2FA | 问题")
            for plan in plans:
                item, workspace = plan["item"], plan["workspace"]
                team = f"#{workspace.id} · {workspace.custom_name or workspace.name or workspace.official_name or '未命名'}" if workspace else "—"
                print(" | ".join(str(value) for value in (
                    item["row"], item["account"], team, item.get("payment_card") or "—", item.get("cost") or "—",
                    item.get("currency") or "—", f"{item.get('expires_on') or '—'} / {plan['db_expiry'] or '—'}",
                    "有" if item.get("totp_secret") else "无", "；".join(plan["issues"]) or "—",
                )))
            ready = [plan for plan in plans if plan["workspace"] is not None and not plan["blocking"]]
            print(f"\n共 {len(plans)} 行：可写入 {len(ready)} 行，跳过 {len(plans) - len(ready)} 行。")
            if not commit:
                print("干跑：没有写库。确认无误后加 --commit 写入。")
                return
            written = expiry_filled = 0
            for plan in ready:
                workspace_id = plan["workspace"].id
                result = await update_workspace_purchase(db, workspace_id, plan["fields"])
                if not result["ok"]:
                    print(f"第 {plan['item']['row']} 行写入失败：{result['error']}")
                    continue
                written += 1
                if plan["item"].get("expires_on") and not plan["db_expiry"]:
                    await update_workspace_expiry(db, workspace_id, date.fromisoformat(plan["item"]["expires_on"]))
                    expiry_filled += 1
            print(f"已写入 {written} 个团队的购买信息，补填到期日 {expiry_filled} 个。")
    finally:
        await engine.dispose()


def main() -> None:
    parser = argparse.ArgumentParser(description="导入 Excel 购买台账到团队购买信息（一次性）")
    sub = parser.add_subparsers(dest="command", required=True)
    ex = sub.add_parser("extract", help="本机：读 Excel，输出 JSON")
    ex.add_argument("xlsx", type=Path)
    ex.add_argument("--out", type=Path, default=Path(".tmp/purchase-import.json"))
    ap = sub.add_parser("apply", help="容器内：按 JSON 干跑或写入")
    ap.add_argument("json", type=Path)
    ap.add_argument("--commit", action="store_true", help="真正写库；不加就是干跑")
    ap.add_argument("--overwrite", action="store_true", help="覆盖已有购买信息")
    ap.add_argument("--assign", action="append", default=[], metavar="行号=团队ID",
                    help="手动指定 Excel 行对应的团队，可重复，例如 --assign 17=24")
    args = parser.parse_args()
    if args.command == "extract":
        extract(args.xlsx, args.out)
    else:
        assign = {}
        for pair in args.assign:
            row, _, workspace_id = pair.partition("=")
            if not (row.strip().isdigit() and workspace_id.strip().isdigit()):
                raise SystemExit(f"--assign 格式应为 行号=团队ID：{pair}")
            assign[int(row)] = int(workspace_id)
        asyncio.run(apply(args.json, args.commit, args.overwrite, assign))


if __name__ == "__main__":
    main()
